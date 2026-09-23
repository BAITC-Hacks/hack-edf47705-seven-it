"""Bounded XLSX import and formula-safe exports for user supplied text."""
import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from forecast.csv_import import CsvImportError, parse_csv


def xlsx_rows(data):
    if len(data) > 5 * 1024 * 1024:
        raise CsvImportError('Файл больше 5 МБ.')
    book = None
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024 or len(archive.infolist()) > 1000:
                raise CsvImportError('Книга слишком большая. Оставьте один лист с историей продаж.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        sheet.reset_dimensions()
        rows = []
        for index, row in enumerate(sheet.iter_rows(max_col=21), start=1):
            if index > 100001:
                raise CsvImportError('В Excel допускается не больше 100 000 строк.')
            if row[-1].value is not None:
                raise CsvImportError('Используйте таблицу не шире 20 колонок.')
            values = []
            for cell in row[:-1]:
                if cell.data_type == 'f':
                    raise CsvImportError('Замените формулы в Excel готовыми значениями перед загрузкой.')
                value = cell.value
                if isinstance(value, (date, datetime)):
                    value = value.strftime('%Y-%m-%d')
                values.append('' if value is None else str(value))
            rows.append(values)
        return rows
    except CsvImportError:
        raise
    except (BadZipFile, KeyError, ValueError, OSError, IndexError, TypeError, ParseError, DefusedXmlException, InvalidFileException) as error:
        raise CsvImportError('Не удалось прочитать Excel. Сохраните файл в формате .xlsx.') from error
    finally:
        if book:
            book.close()


def parse_sales_xlsx(data):
    output = StringIO()
    csv.writer(output).writerows(xlsx_rows(data))
    return parse_csv(output.getvalue().encode('utf-8'))


def workbook_bytes(title, headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    for row_index, values in enumerate([headers, *rows], start=1):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column_index)
            cell.value = ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value
            if isinstance(value, str):
                cell.data_type = 's'  # Text such as =HYPERLINK must never become a formula.
            if isinstance(value, (float, Decimal)):
                cell.number_format = '#,##0.00'
            if isinstance(value, date):
                cell.number_format = 'yyyy-mm-dd'
            if row_index == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1264D6')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 23 if column < len(headers) else 44
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    book.save(output)
    return output.getvalue()
