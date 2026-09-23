from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from forecast.models import Business, Dataset
from forecast.csv_import import CsvImportError, parse_csv
from market.excel import parse_sales_xlsx, workbook_bytes
from market.insights import product_insights
from market.models import Favorite, LedgerEntry, Observation, Product, Profile, Review, SupportMessage


class MarketFixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.seller = get_user_model().objects.create_user('seller')
        cls.buyer = get_user_model().objects.create_user('buyer')
        cls.other = get_user_model().objects.create_user('other')
        Profile.objects.create(user=cls.seller, role='seller')
        Profile.objects.create(user=cls.buyer, role='buyer')
        cls.product = Product.objects.create(seller=cls.seller, title='Кофе', kind='food', description='Кофе в зерне', region='Алматы', price=Decimal('1000'), stock=50)
        cls.today = timezone.localdate()


class MarketTests(MarketFixture):
    def test_login_registration_and_seller_forms_render(self):
        for route in ('market:login', 'market:register'):
            self.assertEqual(self.client.get(reverse(route)).status_code, 200)
        self.client.force_login(self.seller)
        for route in ('market:seller', 'market:product_create'):
            self.assertEqual(self.client.get(reverse(route)).status_code, 200)
        for route in ('market:product_edit', 'market:observation', 'market:history_import'):
            self.assertEqual(self.client.get(reverse(route, args=[self.product.pk])).status_code, 200)

    def test_history_excel_upsert_and_validation_are_atomic(self):
        self.client.force_login(self.seller)
        url = reverse('market:history_import', args=[self.product.pk])
        for sold in (5, 10):
            data = workbook_bytes('История', ['дата', 'цена', 'продано'], [(self.today, 1000, sold)])
            response = self.client.post(url, {'file': SimpleUploadedFile('history.xlsx', data)})
            self.assertEqual(response.status_code, 302)
        self.assertEqual(Observation.objects.count(), 1)
        self.assertEqual(Observation.objects.get().sold, 10)
        content = b'date,price,sold\n2020-01-01,2000,9\n2020-01-02,-1,8\n'
        response = self.client.post(url, {'file': SimpleUploadedFile('bad.csv', content)})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['form'].errors)
        self.assertEqual(Observation.objects.count(), 1)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_ledger_cannot_edit_another_users_entry(self):
        row = LedgerEntry.objects.create(owner=self.buyer, date=self.today, direction='income', amount=12, note='Private')
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('market:ledger_edit', args=[row.pk])).status_code, 404)
        self.client.force_login(self.buyer)
        self.assertEqual(self.client.get(reverse('market:ledger_edit', args=[row.pk])).status_code, 200)

    def test_catalog_and_detail_render(self):
        response = self.client.get(reverse('market:catalog'))
        self.assertContains(response, 'Кофе')
        response = self.client.get(reverse('market:product', args=[self.product.pk]))
        self.assertContains(response, 'Недостаточно данных')
        self.assertContains(response, 'data-set-theme="dark"')

    def test_private_products_not_visible_to_other_users(self):
        self.product.is_published = False
        self.product.save()
        self.assertNotContains(self.client.get(reverse('market:catalog')), 'Кофе')
        for route in ('market:product', 'market:product_excel'):
            self.assertEqual(self.client.get(reverse(route, args=[self.product.pk])).status_code, 404)
        self.client.force_login(self.seller)
        self.assertEqual(self.client.get(reverse('market:product', args=[self.product.pk])).status_code, 200)

    def test_registration_sets_role_without_admin_privileges(self):
        response = self.client.post(reverse('market:register'), {
            'username': 'new-seller', 'role': 'seller', 'password1': 'Complex-flow-492!',
            'password2': 'Complex-flow-492!', 'is_superuser': '1', 'is_staff': '1',
        })
        self.assertRedirects(response, reverse('market:catalog'))
        user = get_user_model().objects.get(username='new-seller')
        self.assertEqual(user.market_profile.role, 'seller')
        self.assertFalse(user.is_staff)
        self.assertFalse(user.is_superuser)

    def test_buyer_cannot_create_and_other_seller_cannot_edit_product(self):
        self.client.force_login(self.buyer)
        self.assertEqual(self.client.post(reverse('market:product_create'), {}).status_code, 403)
        Profile.objects.create(user=self.other, role='seller')
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('market:product_edit', args=[self.product.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse('market:observation', args=[self.product.pk]), {}).status_code, 404)

    def test_seller_creates_product_with_own_identity(self):
        self.client.force_login(self.seller)
        response = self.client.post(reverse('market:product_create'), {
            'title': 'Рис', 'kind': 'food', 'description': 'Мешок', 'region': 'Алматы',
            'price': '12000', 'stock': '8', 'unit': 'мешок', 'lead_days': '7', 'is_published': 'on',
            'seller': self.buyer.pk, 'is_demo': 'on',
        })
        self.assertEqual(response.status_code, 302)
        product = Product.objects.get(title='Рис')
        self.assertEqual(product.seller, self.seller)
        self.assertFalse(product.is_demo)

    def test_observation_upsert_and_future_date(self):
        self.client.force_login(self.seller)
        url = reverse('market:observation', args=[self.product.pk])
        for sold in (5, 9):
            response = self.client.post(url, {'date': self.today, 'price': '1200', 'sold': sold})
            self.assertEqual(response.status_code, 302)
        self.assertEqual(Observation.objects.count(), 1)
        self.assertEqual(Observation.objects.get().sold, 9)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, Decimal('1200'))
        response = self.client.post(url, {'date': self.today + timedelta(days=1), 'price': '1200', 'sold': 2})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['form'].errors)

    def test_review_is_unique_escaped_and_self_review_denied(self):
        url = reverse('market:product', args=[self.product.pk])
        self.client.force_login(self.seller)
        self.assertEqual(self.client.post(url, {'rating': 5, 'comment': 'Мой товар'}).status_code, 403)
        self.client.force_login(self.buyer)
        for rating in (4, 2):
            self.assertEqual(self.client.post(url, {'rating': rating, 'comment': '<script>bad()</script>'}).status_code, 302)
        self.assertEqual(Review.objects.count(), 1)
        self.assertEqual(Review.objects.get().rating, 2)
        self.assertContains(self.client.get(url), '&lt;script&gt;bad()&lt;/script&gt;')
        self.client.post(url, {'rating': 6, 'comment': 'X'})
        self.assertEqual(Review.objects.get().rating, 2)

    def test_hidden_reviews_do_not_affect_rating(self):
        Review.objects.create(product=self.product, author=self.buyer, rating=1, comment='Hidden', is_hidden=True)
        self.assertIsNone(product_insights(self.product)['rating'])
        self.assertNotContains(self.client.get(reverse('market:product', args=[self.product.pk])), 'Hidden')

    def test_favorite_requires_post_and_stays_private(self):
        self.client.force_login(self.buyer)
        url = reverse('market:favorite', args=[self.product.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.client.post(url)
        self.assertEqual(Favorite.objects.count(), 1)
        self.assertContains(self.client.get(reverse('market:catalog'), {'saved': '1'}), 'Кофе')
        self.client.force_login(self.other)
        self.assertNotContains(self.client.get(reverse('market:catalog'), {'saved': '1'}), 'Кофе')

    def test_ledger_exact_totals_owner_isolation_archive_and_restore(self):
        self.client.force_login(self.buyer)
        for direction, amount in [('income', '.30'), ('expense', '.10')]:
            response = self.client.post(reverse('market:ledger'), {'direction': direction, 'amount': amount,
                'date': self.today, 'category': 'other', 'note': 'My entry', 'owner': self.other.pk})
            self.assertEqual(response.status_code, 302)
        LedgerEntry.objects.create(owner=self.other, date=self.today, direction='income', amount=999, note='Private other')
        response = self.client.get(reverse('market:ledger'))
        self.assertEqual(response.context['balance'], Decimal('.20'))
        self.assertNotContains(response, 'Private other')
        foreign = LedgerEntry.objects.get(owner=self.other)
        self.assertEqual(self.client.post(reverse('market:ledger_archive', args=[foreign.pk])).status_code, 404)
        own = LedgerEntry.objects.filter(owner=self.buyer, direction='expense').get()
        url = reverse('market:ledger_archive', args=[own.pk])
        self.client.post(url)
        self.assertEqual(self.client.get(reverse('market:ledger')).context['balance'], Decimal('.30'))
        self.client.post(url)
        self.assertEqual(self.client.get(reverse('market:ledger')).context['balance'], Decimal('.20'))

    def test_ledger_export_uses_only_owner_and_validates_period(self):
        self.client.force_login(self.buyer)
        LedgerEntry.objects.create(owner=self.buyer, date=self.today, direction='income', amount=12, note='=HYPERLINK("bad")')
        LedgerEntry.objects.create(owner=self.other, date=self.today, direction='income', amount=9000, note='SECRET')
        response = self.client.get(reverse('market:ledger_excel'))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['E2'].data_type, 's')
        self.assertEqual(self.client.get(reverse('market:ledger_excel'), {'start': 'invalid'}).status_code, 400)

    @mock.patch.dict('os.environ', {'ANTHROPIC_API_KEY': '', 'ANTHROPIC_AUTH_TOKEN': ''})
    def test_support_fallback_private_history_and_rate_limit(self):
        self.client.force_login(self.buyer)
        url = reverse('market:support')
        for _ in range(5):
            self.assertEqual(self.client.post(url, {'message': 'Как считать расходы?'}).status_code, 302)
        self.assertEqual(self.client.post(url, {'message': 'Ещё вопрос'}).status_code, 429)
        self.assertEqual(SupportMessage.objects.filter(role='assistant', source='faq').count(), 5)
        self.client.force_login(self.other)
        self.assertNotContains(self.client.get(url), 'Как считать расходы?')

    @mock.patch.dict('os.environ', {'ANTHROPIC_API_KEY': 'test'})
    def test_support_ai_response_and_api_failure_fallback(self):
        from market.support import support_answer
        with mock.patch('market.support.anthropic.Anthropic') as client:
            client.return_value.messages.create.return_value = mock.Mock(
                content=[mock.Mock(type='text', text='Откройте Мои финансы.')], stop_reason='end_turn')
            self.assertEqual(support_answer('Где финансы?', []), ('Откройте Мои финансы.', 'ai'))
            client.return_value.messages.create.side_effect = ValueError('Invalid response')
            self.assertEqual(support_answer('Где финансы?', [])[1], 'faq')

    def test_private_forecast_is_not_visible_to_other_session(self):
        file = SimpleUploadedFile('history.csv', b'date,category,quantity\n2026-01-01,Coffee,2\n')
        response = self.client.post(reverse('upload'), {'name': 'Private store', 'business_type': 'retail',
            'region': 'A', 'lead_time_weeks': 2, 'file': file})
        self.assertEqual(response.status_code, 302)
        dataset = Dataset.objects.get()
        url = reverse('dashboard', args=[dataset.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        other_client = Client()
        self.assertEqual(other_client.get(url).status_code, 404)
        self.assertEqual(other_client.post(reverse('recommend', args=[dataset.pk])).status_code, 404)
        self.assertNotContains(other_client.get(reverse('index')), 'Private store')

    def test_authenticated_forecast_owner_can_return_in_new_session(self):
        self.client.force_login(self.buyer)
        business = Business.objects.create(owner=self.buyer, name='Personal', region='Almaty')
        dataset = Dataset.objects.create(business=business, name='Personal history')
        self.assertEqual(self.client.get(reverse('dashboard', args=[dataset.pk])).status_code, 200)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('dashboard', args=[dataset.pk])).status_code, 404)


class InsightTests(MarketFixture):
    def test_price_forecast_stock_and_demand(self):
        Observation.objects.bulk_create([Observation(product=self.product, date=self.today - timedelta(days=ago),
            price=Decimal(1000 - ago * 2), sold=10 if ago < 30 else 5) for ago in range(60)])
        result = product_insights(self.product)
        self.assertEqual(result['trend'], 100)
        self.assertEqual(result['forecast_price'], Decimal('1060.00'))
        self.assertEqual(result['days_left'], 5)
        self.assertTrue(result['restock_soon'])

    def test_sparse_stale_or_single_price_does_not_invent_forecast(self):
        Observation.objects.create(product=self.product, date=self.today, price=1000, sold=100)
        result = product_insights(self.product)
        self.assertIsNone(result['trend'])
        self.assertIsNone(result['forecast_price'])
        self.assertIsNone(result['days_left'])
        Observation.objects.all().update(date=self.today - timedelta(days=50))
        self.assertTrue(product_insights(self.product)['stale'])


class ExcelTests(TestCase):
    def test_demo_catalog_is_repeatable_and_clearly_marked(self):
        from django.core.management import call_command
        from io import StringIO
        call_command('seed_market', stdout=StringIO())
        call_command('seed_market', stdout=StringIO())
        self.assertEqual(Product.objects.filter(is_demo=True).count(), 6)
        self.assertEqual(Observation.objects.count(), 360)
        self.assertEqual(get_user_model().objects.count(), 0)
        response = self.client.get(reverse('market:catalog'))
        self.assertContains(response, 'вымышленные данные')

    def test_xlsx_dates_and_safe_formula_rejection(self):
        data = workbook_bytes('Продажи', ['дата', 'категория', 'количество'], [(timezone.localdate(), 'Кофе', 20)])
        self.assertEqual(parse_sales_xlsx(data).rows, [(timezone.localdate(), 'Кофе', 20)])
        book = Workbook()
        book.active.append(['date', 'category', 'quantity'])
        book.active.append(['2026-01-01', 'Coffee', '=1+1'])
        buf = BytesIO()
        book.save(buf)
        with self.assertRaisesMessage(CsvImportError, 'формулы'):
            parse_sales_xlsx(buf.getvalue())

    def test_invalid_workbook_is_validation_error(self):
        with self.assertRaises(CsvImportError):
            parse_sales_xlsx(b'not an Excel file')

    def test_non_finite_csv_quantity_rejected(self):
        for value in ('nan', 'inf', '-inf'):
            with self.subTest(value=value), self.assertRaises(CsvImportError):
                parse_csv(f'date,category,quantity\n2026-01-01,Coffee,{value}\n'.encode())

    def test_upload_xlsx_end_to_end(self):
        content = workbook_bytes('Продажи', ['date', 'category', 'quantity'], [(timezone.localdate(), 'Coffee', 10)])
        response = self.client.post(reverse('upload'), {'name': 'Excel store', 'business_type': 'retail',
            'region': 'A', 'lead_time_weeks': 2, 'file': SimpleUploadedFile('sales.xlsx', content)})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Dataset.objects.get().records.count(), 1)
